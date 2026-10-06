import pytest

from clock import config as C, files
from clock.tools import run_tool

YES, NO = (lambda q: True), (lambda q: False)


@pytest.fixture
def area(tmp_path, monkeypatch):
    docs, desk = tmp_path / "Documents", tmp_path / "Desktop"
    docs.mkdir()
    desk.mkdir()
    monkeypatch.setattr(C, "ALLOWED_DIR", docs.resolve())
    monkeypatch.setattr(C, "EXTRA_DIRS", [desk.resolve()])
    return docs.resolve(), desk.resolve(), tmp_path


def test_resolve_stays_in_approved_folders(area):
    docs, desk, tmp = area
    assert files.resolve("a.txt") == docs / "a.txt"
    assert files.resolve("Desktop/b.txt") == desk / "b.txt"   # first part names another approved folder
    assert files.resolve(str(desk / "c.txt")) == desk / "c.txt"
    assert files.resolve("../secret.txt") is None
    assert files.resolve(str(tmp / "other.txt")) is None
    assert files.resolve("") is None


def test_folder_link_cannot_escape(area):
    docs, _, tmp = area
    (tmp / "outside").mkdir()
    link = docs / "out"
    try:
        link.symlink_to(tmp / "outside", target_is_directory=True)
    except OSError:
        pytest.skip("symlinks need privileges on this Windows")
    assert files.resolve("out/x.txt") is None


def test_write_create_append_overwrite(area):
    docs, _, _ = area
    assert files.write("n.txt", "one", "create", YES) == "Wrote n.txt."
    assert "already exists" in files.write("n.txt", "two", "create", YES)
    files.write("n.txt", "+two", "append", YES)
    assert (docs / "n.txt").read_text() == "one+two"
    assert files.write("n.txt", "new", "overwrite", NO) == "User declined."
    assert (docs / "n.txt").read_text() == "one+two"
    files.write("n.txt", "new", "overwrite", YES)
    assert (docs / "n.txt").read_text() == "new"


def test_program_files_need_a_yes(area):
    docs, _, _ = area
    assert files.write("run.ps1", "echo hi", "create", NO) == "User declined."
    assert not (docs / "run.ps1").exists()
    assert files.write("run.ps1", "echo hi", "create", YES) == "Wrote run.ps1."


def test_write_outside_is_refused(area):
    _, _, tmp = area
    assert files.write(str(tmp / "evil.txt"), "x", "create", YES) == files.OUTSIDE
    assert not (tmp / "evil.txt").exists()


def test_move_and_copy_confirm_before_replacing(area):
    docs, desk, _ = area
    (docs / "a.txt").write_text("A")
    (desk / "a.txt").write_text("OLD")
    assert files.copy("a.txt", "Desktop/a.txt", NO) == "User declined."
    assert (desk / "a.txt").read_text() == "OLD"
    assert files.copy("a.txt", "Desktop/a.txt", YES) == "Copied a.txt."
    assert (desk / "a.txt").read_text() == "A"
    assert files.move("a.txt", "Desktop/moved.txt", NO) == "Moved a.txt."  # nothing to replace: no question
    assert not (docs / "a.txt").exists() and (desk / "moved.txt").read_text() == "A"
    assert files.move("moved.txt", "../x.txt", YES) == files.OUTSIDE


def test_delete_always_asks_and_never_touches_roots(area, monkeypatch):
    docs, _, _ = area
    (docs / "d.txt").write_text("x")
    recycled = []
    monkeypatch.setattr(files, "_recycle", lambda p: recycled.append(p.name) or True)
    assert files.delete("d.txt", NO) == "User declined." and not recycled
    assert "Recycle Bin" in files.delete("d.txt", YES) and recycled == ["d.txt"]
    assert "approved folder itself" in files.delete(str(docs), YES)
    assert files.delete("../x", YES) == files.OUTSIDE


def test_list_and_search_cover_every_approved_folder(area):
    docs, desk, _ = area
    (docs / "report cypher.txt").write_text("1")
    (desk / "cypher notes.txt").write_text("2")
    (docs / "sub").mkdir()
    assert "sub/" in files.list_folder("") or "Approved folders" in files.list_folder("")
    assert "sub/" in files.list_folder("Documents")
    hits = files.search("cypher")
    assert "report cypher.txt" in hits and "cypher notes.txt" in hits


def test_tools_are_registered_with_confirm_hidden_from_the_model():
    from clock.tools import TOOLS
    by = {t["name"]: t["input_schema"] for t in TOOLS}
    for name in ("list_folder", "write_file", "make_folder", "move_file", "copy_file", "delete_file"):
        assert name in by
        assert "confirm" not in by[name]["properties"]
    assert by["write_file"]["properties"]["mode"]["enum"] == ["create", "append", "overwrite"]


def test_run_tool_passes_confirm(area):
    docs, _, _ = area
    (docs / "k.txt").write_text("1")
    assert run_tool("write_file", {"path": "k.txt", "text": "2", "mode": "overwrite"}, None, NO) == "User declined."


def test_real_recycle_bin_roundtrip(tmp_path):
    import sys
    if sys.platform != "win32":
        pytest.skip("Windows only")
    f = tmp_path / "bin_me.txt"
    f.write_text("x")
    assert files._recycle(f) is True
    assert not f.exists()


def _pdf_with_text(text):
    """A tiny valid one-page PDF whose page really holds `text` (no PDF writer library needed)."""
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % off for off in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)


def test_reads_word_documents_with_tables(area):
    import docx
    docs, _, _ = area
    d = docx.Document()
    d.add_paragraph("Assignment 3: sorting algorithms")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Deadline", "Friday"
    d.save(docs / "a3.docx")
    out = files.read("a3.docx")
    assert "sorting algorithms" in out and "Deadline | Friday" in out


def test_reads_pdf_text(area):
    docs, _, _ = area
    (docs / "notes.pdf").write_bytes(_pdf_with_text("Quarterly results are up"))
    assert "Quarterly results are up" in files.read("notes.pdf")


def test_broken_or_scanned_files_give_a_spoken_reason(area):
    docs, _, _ = area
    (docs / "bad.docx").write_text("not a zip")
    (docs / "bad.pdf").write_text("not a pdf")
    (docs / "blank.pdf").write_bytes(_pdf_with_text(""))
    assert "couldn't open that Word file" in files.read("bad.docx")
    assert "couldn't open that PDF" in files.read("bad.pdf")
    assert "no text I can read" in files.read("blank.pdf")


def test_long_files_are_read_in_windows(area, monkeypatch):
    docs, _, _ = area
    monkeypatch.setattr(files, "MAX_READ", 10)
    (docs / "long.txt").write_text("0123456789abcdefghij")
    first = files.read("long.txt")
    assert first.startswith("0123456789") and "offset 10 of 20" in first
    assert files.read("long.txt", 10) == "abcdefghij"
    assert "past the end" in files.read("long.txt", 99)
